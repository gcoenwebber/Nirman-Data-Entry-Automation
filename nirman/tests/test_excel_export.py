"""
Tests for the NIRMAN Excel export module.
"""

import os
import pytest
from nirman.models import (
    FieldDefinition, FieldType, FormSchema, FieldValue,
    ExtractionRecord, ExtractionStatus, RecordStatus,
    BatchAnalytics,
)
from nirman.analytics import compute_analytics
from nirman.excel_export import export_workbook


@pytest.fixture
def sample_schema():
    return FormSchema(
        name="Test Form",
        fields=[
            FieldDefinition(name="Q1", label="Content", field_type=FieldType.RATING, min_value=1, max_value=5),
            FieldDefinition(name="Q2", label="Trainer", field_type=FieldType.RATING, min_value=1, max_value=5),
        ]
    )


def make_record(q1_val, q2_val, status=RecordStatus.SUCCESS, filename="test.jpg"):
    rec = ExtractionRecord(schema_id="test", source_file=filename, status=status)
    rec.fields["Q1"] = FieldValue(raw_value=q1_val, validated_value=q1_val, status=ExtractionStatus.VALID)
    rec.fields["Q2"] = FieldValue(raw_value=q2_val, validated_value=q2_val, status=ExtractionStatus.VALID)
    return rec


class TestExcelExport:
    def test_export_creates_file(self, sample_schema, tmp_path):
        records = [make_record(4, 5, filename="form1.jpg")]
        analytics = compute_analytics(records, sample_schema)
        output = str(tmp_path / "test_output.xlsx")

        result = export_workbook(records, sample_schema, analytics, output)
        assert os.path.isfile(result)

    def test_export_multiple_records(self, sample_schema, tmp_path):
        records = [
            make_record(4, 5, filename="form1.jpg"),
            make_record(3, 4, filename="form2.jpg"),
            make_record(5, 3, filename="form3.jpg"),
        ]
        analytics = compute_analytics(records, sample_schema)
        output = str(tmp_path / "test_output.xlsx")

        export_workbook(records, sample_schema, analytics, output)
        assert os.path.isfile(output)

        # Verify sheets exist
        from openpyxl import load_workbook
        wb = load_workbook(output)
        assert "Data Records" in wb.sheetnames
        assert "Field Definitions" in wb.sheetnames
        assert "Data Quality" in wb.sheetnames
        assert "Summary Statistics" in wb.sheetnames

        # Verify data records has correct row count
        ws = wb["Data Records"]
        assert ws.max_row == 4  # 1 header + 3 data rows

    def test_export_with_sentinel_values(self, sample_schema, tmp_path):
        rec = ExtractionRecord(schema_id="test", source_file="form.jpg", status=RecordStatus.PARTIAL)
        rec.fields["Q1"] = FieldValue(
            raw_value="ILLEGIBLE", validated_value="ILLEGIBLE",
            status=ExtractionStatus.ILLEGIBLE, needs_review=True
        )
        rec.fields["Q2"] = FieldValue(raw_value=4, validated_value=4, status=ExtractionStatus.VALID)
        rec.review_required = True

        analytics = compute_analytics([rec], sample_schema)
        output = str(tmp_path / "test_output.xlsx")
        export_workbook([rec], sample_schema, analytics, output)
        assert os.path.isfile(output)

    def test_export_empty_records(self, sample_schema, tmp_path):
        analytics = compute_analytics([], sample_schema)
        output = str(tmp_path / "test_output.xlsx")
        export_workbook([], sample_schema, analytics, output)
        assert os.path.isfile(output)
